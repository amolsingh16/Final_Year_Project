import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

print("Building Document Generator Script...")

# Document Output Paths
DOCX_PATH = os.path.abspath("Krishi_Mitra_Project_Complete_Explanation.docx")
PDF_PATH = os.path.abspath("Krishi_Mitra_Project_Complete_Explanation.pdf")

# Color Palette Definitions
PRIMARY_COLOR = RGBColor(27, 67, 50)      # #1B4332 - Dark Green
SECONDARY_COLOR = RGBColor(45, 106, 79)   # #2D6A4F - Medium Green
DARK_TEXT = RGBColor(33, 37, 41)         # #212529 - Charcoal Text
LIGHT_BG = RGBColor(248, 249, 250)       # #F8F9FA - Light Gray

HEX_PRIMARY = "#1B4332"
HEX_SECONDARY = "#2D6A4F"
HEX_LIGHT_BG = "#F8F9FA"

# ---------------------------------------------------------
# 1. HELPER FUNCTIONS FOR DOCX
# ---------------------------------------------------------

def set_cell_background(cell, fill_color_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "E8F5E9") # Light green callout bg
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border green
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="1B4332"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}: ")
    r_title.bold = True
    r_title.font.color.rgb = PRIMARY_COLOR
    r_text = p.add_run(text)
    r_text.font.color.rgb = DARK_TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_str):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "222830") # Dark code bg
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(code_str)
    r.font.name = 'Consolas'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(233, 236, 239) # Light text
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ---------------------------------------------------------
# 2. REPORTLAB PDF NUMBERED CANVAS
# ---------------------------------------------------------

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        if self._pageNumber == 1:
            # Suppress header/footer on title page
            return
        
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#6C757D"))
        
        # Top Header
        self.drawString(54, 750, "Krishi Mitra – B.Tech Final Year Viva Preparation Guide")
        self.setStrokeColor(colors.HexColor("#DEE2E6"))
        self.setLineWidth(0.5)
        self.line(54, 742, 558, 742)
        
        # Bottom Footer
        self.line(54, 45, 558, 45)
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 32, page_text)
        self.drawString(54, 32, "Faculty Guide: Mohd. Osama (PhD) | Author: Amol Singh")
        self.restoreState()

print("Helper classes defined. Generating document data...")
